#!/usr/bin/env python3
"""Render sentence-level subtitles for an existing approval-copy document."""

from __future__ import annotations

import argparse
import json
import os
import re
from pathlib import Path
from enum import Enum
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from article_paragraphs import render as render_paragraphs
from subtitle_timing import (
    SUBTITLE_MAX_WIDTH_UNITS,
    display_text,
    display_width_units,
    split_timeline_segments,
    visible_subtitle_segments,
)
from video_constants import FOOTER_CAPTIONS


SHOT_HEADING = re.compile(r"^#{2,3}\s*镜头\s*(\d+)(?:\s*[（(].*?[）)])?\s*$", re.MULTILINE)


def atomic_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(content, encoding="utf-8", newline="\n")
    os.replace(temporary, path)


def require_unapproved_copy(job_path: Path, removed_shots: list[int]) -> None:
    """Refuse to rewrite copy a human has already approved.

    Removing a duplicate shot changes approved wording, so it needs a fresh
    human approval and a new seal; it must never re-sign the snapshot itself.
    """
    if not job_path.is_file():
        return
    job = json.loads(job_path.read_text(encoding="utf-8-sig"))
    approval = job.get("copy_approval") if isinstance(job, dict) else None
    if isinstance(approval, dict) and approval.get("status") == "approved":
        raise ValueError(
            f"已审批的文案中镜头 {', '.join(map(str, removed_shots))} 与前面的镜头重复；"
            "请修订 copy-for-approval.md 并重新提交人工审批，不能自动改写已审批文案"
        )


# Labelled lines a shot block may carry besides its narration. They are read
# by build_video_job.py and must never become part of the spoken narration.
COPY_FIELD_LABELS = ("证据", "文章依据", "文章证据", "模块", "画面", "衔接", "表题")
NARRATION_METADATA_PREFIXES = (
    *(f"{label}{colon}" for label in COPY_FIELD_LABELS for colon in "：:"),
    "数据表",
    "---",
    "字幕分段",
    "|",
)
NARRATION_MARKUP_PREFIXES = ("- ", "**", "##", "#")


def _join_narration_lines(lines: list[str]) -> str:
    """Join a wrapped Markdown paragraph; Chinese lines join without a space."""
    text = lines[0]
    for line in lines[1:]:
        latin_boundary = text[-1:].isascii() and text[-1:].isalnum() and line[:1].isascii() and line[:1].isalnum()
        text += (" " if latin_boundary else "") + line
    return text


def shot_narrations(copy: str) -> list[tuple[int, str]]:
    """Read each shot's narration: the first paragraph under its heading.

    A paragraph may wrap over several lines; every line up to the next blank
    line or metadata line belongs to the narration.
    """
    matches = list(SHOT_HEADING.finditer(copy))
    shots: list[tuple[int, str]] = []
    for position, match in enumerate(matches):
        end = matches[position + 1].start() if position + 1 < len(matches) else len(copy)
        paragraph: list[str] = []
        for line in copy[match.end():end].splitlines():
            value = line.strip()
            if (
                not value
                or value.startswith(NARRATION_METADATA_PREFIXES)
                or value.startswith(NARRATION_MARKUP_PREFIXES)
            ):
                if paragraph:
                    break
                continue
            if not paragraph:
                value = re.sub(r"^口播[：:]\s*", "", value)
                if not value:
                    continue
            paragraph.append(value)
        if paragraph:
            shots.append((int(match.group(1)), _join_narration_lines(paragraph)))
    if not shots:
        raise ValueError("copy-for-approval.md must contain headings such as '### 镜头1' followed by narration")
    return shots


def validate_generation_subtitles(shots: list[tuple[int, str]]) -> None:
    """Fail before writing approval files only when narration cannot render."""
    for index, narration in shots:
        try:
            visible_subtitle_segments(narration)
        except ValueError as error:
            raise ValueError(
                f"shot {index} subtitle cannot fit the template width: {error}"
            ) from error


# A data_table shot renders at most 5 rows, so a narration that names more
# parameter values than that cannot become one table without dropping content.
TABLE_ROW_LIMIT = 5
PARAMETER_VALUE_PATTERN = re.compile(
    r"\d+(?:\.\d+)?(?:\s*[-~～至]\s*\d+(?:\.\d+)?)?\s*"
    r"(?:%|％|℃|°C|MPa|kPa|GPa|kJ/m²|kJ/m2|g/cm³|g/cm3|g/10min|kg/m³|kg/m3|h|小时|min|分钟|秒|倍|次|级|V|A|W)"
)
PARAMETER_TABLE_HEADERS = {"项目", "参数", "指标", "property", "properties", "item"}


def table_parameter_warnings(shots: list[tuple[int, str]]) -> list[dict[str, Any]]:
    """Warn when a shot names more parameter values than one table can hold."""
    report: list[dict[str, Any]] = []
    for index, narration in shots:
        values = len(PARAMETER_VALUE_PATTERN.findall(narration))
        if values > TABLE_ROW_LIMIT:
            report.append({
                "shot": index,
                "parameters": values,
                "limit": TABLE_ROW_LIMIT,
                "hint": (
                    "a data_table shot renders at most 5 rows; name at most 5 parameters here, "
                    "spread them over an extra table shot (budget permitting), or remove the extra "
                    "value from the narration too so the table and the voice-over stay consistent"
                ),
            })
    return report


def table_capacity_warning(article_text: str) -> str | None:
    """Warn when the article carries more parameter rows than the budget renders."""
    from validate_video_job import DEFAULT_SHOT_BUDGET

    rows = 0
    for line in str(article_text or "").splitlines():
        if "|" not in line:
            continue
        if re.fullmatch(r"[\s|:\-]+", line):
            continue  # a Markdown separator row such as |---|---|
        first_cell = line.strip().strip("|").split("|", 1)[0].strip().strip("*_ ")
        if first_cell.casefold() in PARAMETER_TABLE_HEADERS:
            continue
        rows += 1
    capacity = TABLE_ROW_LIMIT * DEFAULT_SHOT_BUDGET["max_data_tables"]
    if rows <= capacity:
        return None
    return (
        f"the article carries {rows} parameter rows but the default budget renders at most "
        f"{capacity} ({DEFAULT_SHOT_BUDGET['max_data_tables']} tables × {TABLE_ROW_LIMIT} rows); "
        "trim the tabulated rows, or declare a larger shot_budget with a reason"
    )


def shot_budget_warning(shot_count: int) -> str | None:
    """Warn before approval when the copy spends more shots than the budget."""
    from validate_video_job import DEFAULT_SHOT_BUDGET

    limit = DEFAULT_SHOT_BUDGET["max_total"]
    if shot_count <= limit:
        return None
    return (
        f"{shot_count} shots exceeds the default {limit}-shot budget; consolidate the modules "
        "before approval, or declare a larger shot_budget with a reason in video-job.json."
    )


def overlong_clauses(shots: list[tuple[int, str]]) -> list[dict[str, Any]]:
    """List clauses the renderer must wrap so the copy can be rewritten first."""
    report: list[dict[str, Any]] = []
    for index, narration in shots:
        for clause in split_timeline_segments(narration):
            visible = display_text(clause)
            width = display_width_units(visible)
            if width > SUBTITLE_MAX_WIDTH_UNITS:
                report.append({"shot": index, "clause": visible, "width": width})
    return report


def _comparison_text(value: str) -> str:
    """Remove formatting and generic glue words before comparing shot copy."""
    # Drop only the grade token itself (prefix plus its ASCII code). `\w` would
    # also swallow the Chinese clause after it, so "这款PP Z30S耐热" and
    # "这款PP Z30S刚性好" both collapsed to "这款" and were deleted as duplicates.
    value = re.sub(
        r"(?<![A-Za-z])(?:HDPE|LDPE|LLDPE|PC|PP|PE|PVC|ABS|POM|PET|PA66?)(?:\s*[-/]?\s*[A-Za-z0-9][A-Za-z0-9-]*)?",
        "",
        value,
        flags=re.IGNORECASE,
    )
    return re.sub(r"[\s，,、；;：:。！？!?（）()「」『』“”\"'、·]", "", value)


def _char_ngrams(value: str, size: int = 3) -> set[str]:
    return {value[index : index + size] for index in range(max(0, len(value) - size + 1))}


class OverlapKind(str, Enum):
    DISTINCT = "distinct"
    EXACT = "exact"
    SUBSTANTIAL = "substantial"


def classify_overlap(left: str, right: str) -> OverlapKind:
    left_text, right_text = _comparison_text(left), _comparison_text(right)
    if not left_text or not right_text:
        return OverlapKind.DISTINCT
    if left_text == right_text:
        return OverlapKind.EXACT
    shorter, longer = sorted((left_text, right_text), key=len)
    if len(shorter) >= 12 and shorter in longer:
        return OverlapKind.SUBSTANTIAL
    left_ngrams, right_ngrams = _char_ngrams(left_text), _char_ngrams(right_text)
    if not left_ngrams or not right_ngrams:
        return OverlapKind.DISTINCT
    if (
        len(left_ngrams & right_ngrams) / len(left_ngrams | right_ngrams) >= 0.55
        or len(left_ngrams & right_ngrams) / min(len(left_ngrams), len(right_ngrams)) >= 0.78
    ):
        return OverlapKind.SUBSTANTIAL
    return OverlapKind.DISTINCT


def validate_shot_distinctness(shots: list[tuple[int, str]]) -> None:
    """Reject LLM copy that repeats a substantive clause across shots."""
    prepared = [(index, _comparison_text(narration)) for index, narration in shots]
    for left_position, (left_index, left_text) in enumerate(prepared):
        if not left_text:
            continue
        for right_index, right_text in prepared[left_position + 1 :]:
            if not right_text:
                continue
            overlap = classify_overlap(left_text, right_text)
            if overlap is OverlapKind.EXACT:
                raise ValueError(f"镜头{left_index}与镜头{right_index}文案重复，请改写为不同信息点")
            if overlap is OverlapKind.SUBSTANTIAL:
                shorter, longer = sorted((left_text, right_text), key=len)
                if len(shorter) >= 12 and shorter in longer:
                    raise ValueError(f"镜头{left_index}与镜头{right_index}复用了同一段核心文案，请拆分信息点")
                raise ValueError(f"镜头{left_index}与镜头{right_index}内容高度重叠，请确保每个镜头只承载一个独立信息点")


def _shots_overlap(left: str, right: str) -> bool:
    return classify_overlap(left, right) is OverlapKind.EXACT


def remove_duplicate_shot_blocks(copy: str) -> tuple[str, list[tuple[int, str]], list[int]]:
    """Drop later duplicate shot sections so generation can continue automatically."""
    matches = list(SHOT_HEADING.finditer(copy))
    if not matches:
        return copy, [], []
    kept: list[tuple[int, str]] = []
    removed: list[int] = []
    chunks: list[str] = [copy[: matches[0].start()]]
    next_index = 1
    for position, match in enumerate(matches):
        end = matches[position + 1].start() if position + 1 < len(matches) else len(copy)
        block = copy[match.start() : end]
        parsed = shot_narrations(block)
        if not parsed:
            continue
        shot = parsed[0]
        if any(_shots_overlap(shot[1], prior[1]) for prior in kept):
            removed.append(shot[0])
            continue
        kept.append((next_index, shot[1]))
        # Renumber only the heading digits and keep the block's own text and
        # line breaks, so the next heading still starts on its own line.
        chunks.append(re.sub(r"^(#{2,3}\s*镜头\s*)\d+", rf"\g<1>{next_index}", block, count=1))
        next_index += 1
    if not kept:
        raise ValueError("自动去重后没有剩余有效镜头文案")
    return "".join(chunks), kept, removed


def render_subtitles(title: str, product_grade: str, shots: list[tuple[int, str]]) -> str:
    lines = [
        f"# {title or product_grade + ' 视频'} 字幕（待审批）",
        "",
        f"- **产品牌号**：{product_grade}",
        "- 每行是一条将导入剪映时间轴的最终可见字幕；审批内容与成片字幕逐条一致。",
    ]
    for index, narration in shots:
        lines.extend(["", f"### 镜头{index}", ""])
        lines.extend(visible_subtitle_segments(narration))
    lines.extend(["", "### 片尾"])
    for caption in (template.format(product_grade=product_grade) for template in FOOTER_CAPTIONS):
        lines.extend(visible_subtitle_segments(caption))
    return "\n".join(lines) + "\n"


def render_subtitles_docx(
    destination: Path,
    title: str,
    product_grade: str,
    shots: list[tuple[int, str]],
    subtitle_sections: list[tuple[str, list[str]]] | None = None,
) -> None:
    """Create a readable, comment-friendly Word version of the approval subtitles."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(4)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(f"{title or product_grade + ' 视频'} 字幕（待审批）")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 78, 121)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(f"产品牌号：{product_grade}    |    每行对应一条最终可见字幕")
    meta_run.font.name = "Microsoft YaHei"
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(89, 89, 89)

    def add_section(label: str, captions: list[str]) -> None:
        section_heading = document.add_paragraph()
        section_heading.paragraph_format.space_before = Pt(10)
        section_heading.paragraph_format.space_after = Pt(3)
        section_run = section_heading.add_run(label)
        section_run.bold = True
        section_run.font.name = "Microsoft YaHei"
        section_run.font.size = Pt(13)
        section_run.font.color.rgb = RGBColor(31, 78, 121)
        for caption in captions:
            paragraph = document.add_paragraph(style="Normal")
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.keep_together = True
            paragraph.add_run(caption)

    if subtitle_sections is None:
        subtitle_sections = [(f"镜头{index}", visible_subtitle_segments(narration)) for index, narration in shots]
        subtitle_sections.append(("片尾", [segment for caption in (template.format(product_grade=product_grade) for template in FOOTER_CAPTIONS) for segment in visible_subtitle_segments(caption)]))
    for label, captions in subtitle_sections:
        add_section(label, captions)
    document.core_properties.title = f"{title or product_grade + ' 视频'} 字幕（待审批）"
    document.core_properties.subject = "短视频字幕审批稿"
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("article", type=Path, help="article.json source record")
    parser.add_argument("--output-dir", required=True, type=Path, help="Job directory containing copy-for-approval.md")
    parser.add_argument("--product-grade", required=True, help="Complete product grade, for example HDPE 8008")
    parser.add_argument("--copy-for-approval", type=Path, help="Optional approval-copy path")
    args = parser.parse_args()

    article_path = args.article.resolve()
    output_dir = args.output_dir.resolve()
    copy_path = (args.copy_for_approval or output_dir / "copy-for-approval.md").resolve()
    if not article_path.is_file():
        parser.error(f"article.json not found: {article_path}")
    if not copy_path.is_file():
        parser.error(f"approval copy not found: {copy_path}")
    article = json.loads(article_path.read_text(encoding="utf-8-sig"))
    if not isinstance(article, dict):
        parser.error("article.json root must be an object")
    original_copy = copy_path.read_text(encoding="utf-8")
    adjusted_copy, shots, removed_shots = remove_duplicate_shot_blocks(original_copy)
    job_path = output_dir / "video-job.json"
    if removed_shots:
        require_unapproved_copy(job_path, removed_shots)
        atomic_write(copy_path, adjusted_copy)
    validate_generation_subtitles(shots)
    validate_shot_distinctness(shots)
    # The reviewed copy is the source of truth: always re-render the subtitle
    # list so a revised clause cannot leave the two approval files out of sync.
    content = render_subtitles(str(article.get("title") or "").strip(), args.product_grade.strip(), shots)
    destination = output_dir / "subtitle-for-approval.md"
    atomic_write(destination, content)
    docx_destination = output_dir / "subtitle-for-approval.docx"
    render_subtitles_docx(docx_destination, str(article.get("title") or "").strip(), args.product_grade.strip(), shots)
    # Paragraph numbers let `证据：` cite P3 instead of a hand-copied quote.
    atomic_write(output_dir / "article-paragraphs.md", render_paragraphs(str(article.get("body_text") or "")))
    print(json.dumps({
        "output": str(destination),
        "docx_output": str(docx_destination),
        "shots": len(shots),
        "removed_duplicate_shots": removed_shots,
        "overlong_clauses": overlong_clauses(shots),
        "shot_budget_warning": shot_budget_warning(len(shots)),
        "table_parameter_warnings": table_parameter_warnings(shots),
        "table_capacity_warning": table_capacity_warning(str(article.get("body_text") or "")),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
